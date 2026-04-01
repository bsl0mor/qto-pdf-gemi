import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from utils import load_qty_from_result

HERE = os.path.dirname(os.path.abspath(__file__))

# ─── PATHS ────────────────────────────────────────────────────────────────────
PROJECTS = [
    {
        "name":       "3246",
        "sub_json":   r"C:\Users\basel\Downloads\QTO-AI\3246_result.json",
        "super_json": r"C:\Users\basel\3246_super_result.json",
        "arch_json":  r"C:\Users\basel\Downloads\QTO-AI\3246_arch_result.json",
    },
    {
        "name":       "1577",
        "sub_json":   r"C:\Users\basel\Downloads\QTO-AI\1577_result.json",
        "super_json": r"C:\Users\basel\1577_super_result.json",
        "arch_json":  r"C:\Users\basel\Downloads\QTO-AI\1577_arch_result.json",
    },
]

OUT = r"C:\Users\basel\Downloads\QTO-AI\BOQ_Export_v2.xlsx"

# ─── DESCRIPTION MAP ──────────────────────────────────────────────────────────
DESC = {
    "excavation":          "Bulk Excavation",
    "road_base":           "Road Base",
    "pcc_footings":        "PCC Blinding under Isolated Footings",
    "pcc_tb":              "PCC Blinding under Tie Beams",
    "footing_concrete":    "Isolated Footing Concrete",
    "neck_column":         "Neck Column Concrete",
    "tb_concrete":         "Tie Beam RC Concrete",
    "bitumen":             "Bituminous Waterproofing Sub-Structure",
    "blockwall_solid_sub": "Solid Block Sub-Structure",
    "slab_on_grade":       "Slab on Grade Concrete",
    "polythene":           "Polythene Sheet 1000 gauge",
    "anti_termite":        "Anti-Termite Treatment",
    "backfill":            "Backfill Compacted",
    "gf_columns":          "GF Columns Concrete",
    "ff_columns":          "FF Columns Concrete",
    "ff_beams":            "First Floor Beams Concrete",
    "ff_slab":             "First Floor Slab Concrete",
    "roof_beams":          "Roof Beams Concrete",
    "roof_slab":           "Roof Slab Concrete",
    "parapet_concrete":    "Parapet Coping Concrete",
    "parapet_block":       "Parapet Block Work",
    "block_20_ext_gf":     "External Block 20cm — Ground Floor",
    "block_20_int_gf":     "Internal Block 20cm — Ground Floor",
    "block_10_int_gf":     "Internal Block 10cm — Ground Floor",
    "plaster_int_gf":      "Internal Plaster — Ground Floor",
    "flooring_dry_gf":     "Flooring Dry Areas — Ground Floor",
    "flooring_wet_gf":     "Flooring Wet Areas — Ground Floor",
    "skirting_gf":         "Skirting — Ground Floor",
    "paint_gf":            "Paint Internal Walls — Ground Floor",
    "ceiling_dry_gf":      "Ceiling Dry Areas — Ground Floor",
    "ceiling_wet_gf":      "Ceiling Wet Areas — Ground Floor",
    "tiles_wall_gf":       "Wall Tiles Wet Areas — Ground Floor",
    "block_20_ext_ff":     "External Block 20cm — First Floor",
    "block_20_int_ff":     "Internal Block 20cm — First Floor",
    "block_10_int_ff":     "Internal Block 10cm — First Floor",
    "plaster_int_ff":      "Internal Plaster — First Floor",
    "flooring_dry_ff":     "Flooring Dry Areas — First Floor",
    "flooring_wet_ff":     "Flooring Wet Areas — First Floor",
    "skirting_ff":         "Skirting — First Floor",
    "paint_ff":            "Paint Internal Walls — First Floor",
    "ceiling_dry_ff":      "Ceiling Dry Areas — First Floor",
    "ceiling_wet_ff":      "Ceiling Wet Areas — First Floor",
    "tiles_wall_ff":       "Wall Tiles Wet Areas — First Floor",
    "finish_ext":          "External Wall Finish",
    "marble_threshold":    "Marble Threshold",
    "balcony_flooring":    "Balcony Flooring",
    "balcony_wp":          "Balcony Waterproofing",
    "waterproofing":       "Waterproofing FF Wet Areas",
    "roof_waterproofing":  "Roof Waterproofing",
    "combo_roof":          "Roof Combo (Screed + WP + Insulation)",
    "doors_schedule":      "Doors Schedule",
    "windows_schedule":    "Windows Schedule",
}

# ─── SECTIONS ─────────────────────────────────────────────────────────────────
SECTIONS = [
    ("SUBSTRUCTURE", [
        "excavation", "road_base", "pcc_footings", "pcc_tb",
        "footing_concrete", "neck_column", "tb_concrete",
        "bitumen", "blockwall_solid_sub", "slab_on_grade",
        "polythene", "anti_termite", "backfill",
    ]),
    ("SUPERSTRUCTURE", [
        "gf_columns", "ff_columns", "ff_beams", "ff_slab",
        "roof_beams", "roof_slab", "parapet_concrete", "parapet_block",
    ]),
    ("ARCHITECTURAL FINISHES — GROUND FLOOR", [
        "block_20_ext_gf", "block_20_int_gf", "block_10_int_gf",
        "plaster_int_gf", "flooring_dry_gf", "flooring_wet_gf",
        "skirting_gf", "paint_gf", "ceiling_dry_gf", "ceiling_wet_gf",
        "tiles_wall_gf",
    ]),
    ("ARCHITECTURAL FINISHES — FIRST FLOOR", [
        "block_20_ext_ff", "block_20_int_ff", "block_10_int_ff",
        "plaster_int_ff", "flooring_dry_ff", "flooring_wet_ff",
        "skirting_ff", "paint_ff", "ceiling_dry_ff", "ceiling_wet_ff",
        "tiles_wall_ff",
    ]),
    ("ARCHITECTURAL FINISHES — GENERAL", [
        "finish_ext", "marble_threshold", "balcony_flooring", "balcony_wp",
        "waterproofing", "roof_waterproofing", "combo_roof", "doors_schedule", "windows_schedule",
    ]),
]

# ─── STYLES ───────────────────────────────────────────────────────────────────
def thin_border():
    s = Side(style="thin", color="CCCCCC")
    return Border(left=s, right=s, top=s, bottom=s)

def thick_border():
    s = Side(style="medium", color="999999")
    return Border(left=s, right=s, top=s, bottom=s)

COLOR_HEADER    = "1F3864"   # dark navy
COLOR_SECTION   = "2E75B6"   # blue
COLOR_ALT       = "EBF3FB"   # light blue alternate row
COLOR_WHITE     = "FFFFFF"
COLOR_ACCENT    = "F4B942"   # amber for unit column

def hdr_font(size=10, color="FFFFFF", bold=True):
    return Font(name="Calibri", size=size, bold=bold, color=color)

def body_font(bold=False):
    return Font(name="Calibri", size=10, bold=bold)

def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def center():
    return Alignment(horizontal="center", vertical="center", wrap_text=True)

def left():
    return Alignment(horizontal="left", vertical="center", wrap_text=True)

# ─── LOAD ALL QTOSITEMS INTO FLAT DICT ────────────────────────────────────────
def load_qty(json_path):
    """Return dict: item_id -> {qty, unit, breakdown}. Handles both formats."""
    result = {}
    with open(json_path, encoding="utf-8") as f:
        data = json.load(f)
    # arch.py format: {"items": [{"id","u","q","bd"}, ...]}
    # test2.py / super.py format: {"qtoItems": [{"id","unit","totalQty","breakdown"}, ...]}
    items = data.get("qtoItems") or data.get("items", [])
    for item in items:
        iid = str(item.get("id", "")).strip()
        result[iid] = {
            "qty":  item.get("totalQty") if item.get("totalQty") is not None else item.get("q", 0),
            "unit": item.get("unit") or item.get("u", ""),
            "bd":   item.get("breakdown") or item.get("bd", ""),
        }
    return result

# ─── BUILD UNIT MAP (prefer arch, then super, then sub) ───────────────────────
def merge_all(sub, sup, arch):
    merged = {}
    for d in (sub, sup, arch):
        for k, v in d.items():
            merged[k] = v
    return merged

# ─── WRITE ONE SHEET ──────────────────────────────────────────────────────────
def write_sheet(ws, proj_name, merged):
    # Column widths
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 48
    ws.column_dimensions["D"].width = 8
    ws.column_dimensions["E"].width = 14
    ws.column_dimensions["F"].width = 55

    # Row 1 — Project header
    ws.row_dimensions[1].height = 28
    for col in range(1, 7):
        cell = ws.cell(row=1, column=col)
        cell.fill    = fill(COLOR_HEADER)
        cell.font    = hdr_font(14)
        cell.alignment = center()
        cell.border  = thick_border()
    ws.merge_cells("A1:F1")
    ws.cell(row=1, column=1).value = f"BILL OF QUANTITIES — PROJECT {proj_name}"

    # Row 2 — Column headers
    ws.row_dimensions[2].height = 22
    headers = ["#", "ITEM ID", "DESCRIPTION", "UNIT", "QTY", "BREAKDOWN / FORMULA"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col)
        cell.value     = h
        cell.fill      = fill(COLOR_SECTION)
        cell.font      = hdr_font(10)
        cell.alignment = center()
        cell.border    = thick_border()

    row     = 3
    item_no = 1

    for section_name, item_ids in SECTIONS:
        # Section header row
        ws.row_dimensions[row].height = 20
        for col in range(1, 7):
            cell = ws.cell(row=row, column=col)
            cell.fill      = fill(COLOR_SECTION)
            cell.font      = hdr_font(11)
            cell.border    = thick_border()
            cell.alignment = center()
        ws.merge_cells(f"A{row}:F{row}")
        ws.cell(row=row, column=1).value = section_name
        ws.cell(row=row, column=1).alignment = left()
        row += 1

        for iid in item_ids:
            info = merged.get(iid, {})
            qty  = info.get("qty", "—")
            unit = info.get("unit", "")
            bd   = info.get("bd", "")

            # Alternate row color
            bg = COLOR_ALT if item_no % 2 == 0 else COLOR_WHITE

            ws.row_dimensions[row].height = 18

            cells_data = [
                (1, item_no,            center(), False),
                (2, iid,                left(),   False),
                (3, DESC.get(iid, iid), left(),   False),
                (4, unit,               center(), False),
                (5, round(float(qty), 3) if qty != "—" else "—", center(), True),
                (6, bd,                 left(),   False),
            ]

            for col, val, align, bold in cells_data:
                cell = ws.cell(row=row, column=col)
                cell.value     = val
                cell.font      = body_font(bold)
                cell.alignment = align
                cell.border    = thin_border()
                # Unit cell accent
                if col == 4:
                    cell.fill = fill("FFF2CC")
                    cell.font = body_font(bold=True)
                elif col == 5:
                    cell.fill = fill("E2EFDA")
                    cell.font = Font(name="Calibri", size=10, bold=True, color="375623")
                else:
                    cell.fill = fill(bg)

            row     += 1
            item_no += 1

        # Blank separator
        for col in range(1, 7):
            ws.cell(row=row, column=col).fill = fill("F2F2F2")
        row += 1

    print(f"  Sheet '{proj_name}': {item_no-1} items written, {row-1} rows total")

# ─── MAIN ─────────────────────────────────────────────────────────────────────

def write_batch_boq(project_names, output_path):
    """Generate a multi-sheet BOQ Excel from result JSONs."""
    wb = Workbook()
    wb.remove(wb.active)

    all_data = {}
    for name in project_names:
        sub  = load_qty_from_result(os.path.join(HERE, f"{name}_result.json"))
        sup  = load_qty_from_result(os.path.join(HERE, f"{name}_super_result.json"))
        arch = load_qty_from_result(os.path.join(HERE, f"{name}_arch_result.json"))
        merged = merge_all(sub, sup, arch)
        all_data[name] = merged
        print(f"  {name}: sub={len(sub)}, super={len(sup)}, arch={len(arch)}")

        ws = wb.create_sheet(title=f"Project {name}")
        write_sheet(ws, name, merged)

    # Comparison sheet (if 2+ projects)
    if len(project_names) >= 2:
        _write_comparison(wb, project_names, all_data)

    wb.save(output_path)
    print(f"\n  Saved: {output_path}")


def _write_comparison(wb, names, all_data):
    """Write comparison sheet for first two projects."""
    n1, n2 = names[0], names[1]
    ws = wb.create_sheet(title="Comparison")
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 48
    ws.column_dimensions["D"].width = 8
    ws.column_dimensions["E"].width = 14
    ws.column_dimensions["F"].width = 14

    ws.row_dimensions[1].height = 28
    ws.merge_cells("A1:F1")
    for col in range(1, 7):
        cell = ws.cell(row=1, column=col)
        cell.fill = fill(COLOR_HEADER)
        cell.font = hdr_font(13)
        cell.alignment = center()
        cell.border = thick_border()
    ws.cell(row=1, column=1).value = f"QTO COMPARISON — {n1}  vs  {n2}"

    ws.row_dimensions[2].height = 22
    for col, h in enumerate(["#", "ITEM ID", "DESCRIPTION", "UNIT", f"QTY — {n1}", f"QTY — {n2}"], 1):
        cell = ws.cell(row=2, column=col)
        cell.value = h
        cell.fill = fill(COLOR_SECTION)
        cell.font = hdr_font(10)
        cell.alignment = center()
        cell.border = thick_border()

    row = 3
    item_no = 1
    for section_name, item_ids in SECTIONS:
        ws.row_dimensions[row].height = 20
        ws.merge_cells(f"A{row}:F{row}")
        for col in range(1, 7):
            cell = ws.cell(row=row, column=col)
            cell.fill = fill(COLOR_SECTION)
            cell.font = hdr_font(11)
            cell.border = thick_border()
        ws.cell(row=row, column=1).value = section_name
        ws.cell(row=row, column=1).alignment = left()
        row += 1

        for iid in item_ids:
            bg = COLOR_ALT if item_no % 2 == 0 else COLOR_WHITE
            ws.row_dimensions[row].height = 18

            q1 = all_data[n1].get(iid, {})
            q2 = all_data[n2].get(iid, {})
            unit = q1.get("unit", "") or q2.get("unit", "")
            v1 = q1.get("qty", "—")
            v2 = q2.get("qty", "—")

            for col, val, align in [
                (1, item_no,            center()),
                (2, iid,                left()),
                (3, DESC.get(iid, iid), left()),
                (4, unit,               center()),
                (5, round(float(v1), 3) if v1 != "—" else "—", center()),
                (6, round(float(v2), 3) if v2 != "—" else "—", center()),
            ]:
                cell = ws.cell(row=row, column=col)
                cell.value = val
                cell.alignment = align
                cell.border = thin_border()
                if col == 4:
                    cell.fill = fill("FFF2CC")
                    cell.font = body_font(bold=True)
                elif col in (5, 6):
                    cell.fill = fill("E2EFDA")
                    cell.font = Font(name="Calibri", size=10, bold=True, color="375623")
                else:
                    cell.fill = fill(bg)
                    cell.font = body_font()
            row += 1
            item_no += 1

        for col in range(1, 7):
            ws.cell(row=row, column=col).fill = fill("F2F2F2")
        row += 1

    print(f"  Comparison sheet: {item_no-1} items")


if __name__ == "__main__":
    # Legacy standalone mode (original PROJECTS)
    wb = Workbook()
    wb.remove(wb.active)

    for cfg in PROJECTS:
        print(f"\nLoading project {cfg['name']}...")
        sub  = load_qty(cfg["sub_json"])
        sup  = load_qty(cfg["super_json"])
        arch = load_qty(cfg["arch_json"])
        merged = merge_all(sub, sup, arch)
        print(f"  Items loaded: sub={len(sub)}, super={len(sup)}, arch={len(arch)}")

        ws = wb.create_sheet(title=f"Project {cfg['name']}")
        write_sheet(ws, cfg["name"], merged)

    wb.save(OUT)
    print(f"\nSaved: {OUT}")
